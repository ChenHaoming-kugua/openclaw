#!/usr/bin/env python
# -*- coding: utf-8 -*-
"""
读 "码上放心" 导出的 "连锁企业绑定详情" Excel,
过滤 授权状态==是 的行,输出 JSON 到 stdout.
用法: python read_excel.py <file1.xlsx> [<file2.xlsx> ...]
依赖: openpyxl
"""
import json
import sys
from pathlib import Path

try:
    from openpyxl import load_workbook
except ImportError:
    print(json.dumps({"error": "openpyxl not installed. Run: pip install openpyxl"}, ensure_ascii=False))
    sys.exit(2)


COLS = {
    "name":      ["企业名称", "name"],
    "appkey":    ["appkey", "AppKey", "APPKEY"],
    "refEntId":  ["refEntId", "RefEntId", "ref_ent_id"],
    "entId":     ["entId", "EntId", "ent_id"],
    "auth":      ["授权状态", "auth", "authStatus"],
}


def find_col(header_row, candidates):
    for idx, cell in enumerate(header_row):
        v = (cell.value or "").strip() if cell.value else ""
        if v in candidates:
            return idx
    return -1


def read_one(path: Path):
    wb = load_workbook(filename=str(path), read_only=True, data_only=True)
    rows_out = []
    for ws in wb.worksheets:
        rows_iter = ws.iter_rows()
        try:
            header = next(rows_iter)
        except StopIteration:
            continue
        idx = {k: find_col(header, v) for k, v in COLS.items()}
        if idx["name"] < 0 or idx["appkey"] < 0:
            continue
        for r in rows_iter:
            def cell(k):
                i = idx[k]
                if i < 0 or i >= len(r):
                    return ""
                v = r[i].value
                return "" if v is None else str(v).strip()
            auth = cell("auth")
            if auth and auth not in ("是", "Y", "y", "true", "True"):
                continue
            name = cell("name")
            if not name:
                continue
            rows_out.append({
                "name":     name,
                "appkey":   cell("appkey"),
                "refEntId": cell("refEntId"),
                "entId":    cell("entId"),
                "source":   path.name,
            })
    return rows_out


def main():
    if len(sys.argv) < 2:
        print(json.dumps({"error": "usage: read_excel.py <file1.xlsx> ..."}, ensure_ascii=False))
        sys.exit(2)
    all_rows = []
    seen = set()
    for arg in sys.argv[1:]:
        p = Path(arg)
        if not p.exists():
            print(json.dumps({"error": f"file not found: {arg}"}, ensure_ascii=False))
            sys.exit(2)
        for row in read_one(p):
            key = row["name"]
            if key in seen:
                continue
            seen.add(key)
            all_rows.append(row)
    print(json.dumps({"rows": all_rows, "count": len(all_rows)}, ensure_ascii=False))


if __name__ == "__main__":
    main()
